import os
import re
import openpyxl
from io import BytesIO

class WinnerService:
    # Column mapping configurations (case-insensitive)
    HEADER_CANDIDATES = {
        "employee_id": [
            "employee id", "employee_id", "employeeid", "id", "emp id", "empid", 
            "emp_id", "emp code", "empcode", "employee code", "emp no", "empno", 
            "participant id", "user id", "sl. no.", "sl no", "sr no", "sr. no."
        ],
        "identifier_fallback": [
            "email-id", "email id", "email_id", "email", "phone no", "phone_no", 
            "phone", "mobile", "phone number", "mobile no", "mobile number"
        ],
        "name": [
            "name", "participant name", "employee name", "fullname", "full name", 
            "participant", "employee", "member name"
        ],
        "weight": [
            "weight", "weight (kg)", "weight(kg)", "weight (lbs)", "weight(lbs)", 
            "wt", "weight kg", "wt (kg)", "wt(kg)", "weight in kg", "initial weight", "final weight"
        ],
        "waist": [
            "waist", "waist (in)", "waist (cm)", "waist(in)", "waist(cm)", 
            "waist size", "waist (inches)", "waist(inches)", "waist in inches", "initial waist", "final waist", 
            "waist circumference", "waist circumferece"
        ],
        "hip": [
            "hip", "hip (in)", "hip (cm)", "hip(in)", "hip(cm)", 
            "hip size", "hip (inches)", "hip(inches)", "hip in inches", "hips", "initial hip", "final hip", 
            "hip circumference", "hip circumferece"
        ]
    }

    @staticmethod
    def clean_employee_id(val) -> str:
        if val is None:
            return ""
        val_str = str(val).strip()
        # Remove trailing .0 or .00 from floating numbers if read as such
        if re.match(r"^\d+\.0+$", val_str):
            val_str = val_str.split(".")[0]
        # Uppercase to ensure case-insensitive matching across files
        return val_str.upper()

    @staticmethod
    def clean_name(name_val) -> str:
        if name_val is None:
            return ""
        name = str(name_val).strip()
        # Clean trailing merged gender strings if present (e.g. "PavithraFemale" -> "Pavithra")
        if name.lower().endswith("female") and len(name) > 6 and name[-7] != " ":
            name = name[:-6].strip()
        elif name.lower().endswith("male") and len(name) > 4 and name[-5] != " ":
            name = name[:-4].strip()
        return name

    @classmethod
    def parse_and_validate_sheet(cls, file_data: bytes, sheet_label: str) -> dict:
        """
        Parses an Excel sheet from bytes.
        Does NOT return success=False merely because individual participants have missing measurements.
        Only returns success=False for catastrophic file-level failures (unreadable file, empty sheet, missing table headers).
        """
        try:
            wb = openpyxl.load_workbook(BytesIO(file_data), data_only=True)
            ws = wb.active
        except Exception as e:
            return {
                "success": False,
                "file_error": f"Invalid Excel format or file could not be read: {str(e)}"
            }

        if not ws or ws.max_row == 0:
            return {
                "success": False,
                "file_error": "Spreadsheet is empty"
            }

        # Dynamically scan top rows (up to 15) to find the table header row
        header_row_idx = None
        header_indices = {}

        max_scan_row = min(15, ws.max_row)
        for r_idx in range(1, max_scan_row + 1):
            row_vals = [ws.cell(row=r_idx, column=c_idx).value for c_idx in range(1, ws.max_column + 1)]
            temp_indices = {}

            # 1. Resolve Name
            for c_idx, cell_val in enumerate(row_vals):
                if cell_val is not None and str(cell_val).strip().lower() in cls.HEADER_CANDIDATES["name"]:
                    temp_indices["name"] = c_idx
                    break

            # 2. Resolve Weight
            for c_idx, cell_val in enumerate(row_vals):
                if cell_val is not None and str(cell_val).strip().lower() in cls.HEADER_CANDIDATES["weight"]:
                    temp_indices["weight"] = c_idx
                    break

            # 3. Resolve Employee ID or Fallback Identifier (Email/Phone/Name)
            for c_idx, cell_val in enumerate(row_vals):
                if cell_val is not None and str(cell_val).strip().lower() in cls.HEADER_CANDIDATES["employee_id"]:
                    temp_indices["employee_id"] = c_idx
                    break
            if "employee_id" not in temp_indices:
                for c_idx, cell_val in enumerate(row_vals):
                    if cell_val is not None and str(cell_val).strip().lower() in cls.HEADER_CANDIDATES["identifier_fallback"]:
                        temp_indices["employee_id"] = c_idx
                        break
            if "employee_id" not in temp_indices and "name" in temp_indices:
                temp_indices["employee_id"] = temp_indices["name"]

            # 4. Resolve Waist & Hip safely using multi-step priority
            # Step 1: Explicit candidate matching for waist
            for c_idx, cell_val in enumerate(row_vals):
                if cell_val is not None and str(cell_val).strip().lower() in cls.HEADER_CANDIDATES["waist"]:
                    temp_indices["waist"] = c_idx
                    break

            # Step 2: Explicit candidate matching for hip
            for c_idx, cell_val in enumerate(row_vals):
                if cell_val is not None and str(cell_val).strip().lower() in cls.HEADER_CANDIDATES["hip"]:
                    temp_indices["hip"] = c_idx
                    break

            # Step 3: Search for columns containing "waist" if waist is still missing
            if "waist" not in temp_indices:
                for c_idx, cell_val in enumerate(row_vals):
                    if cell_val is not None and "waist" in str(cell_val).strip().lower():
                        temp_indices["waist"] = c_idx
                        break

            # Step 4: Search for columns containing "hip" if hip is still missing
            if "hip" not in temp_indices:
                for c_idx, cell_val in enumerate(row_vals):
                    if cell_val is not None and "hip" in str(cell_val).strip().lower():
                        if c_idx != temp_indices.get("waist"):
                            temp_indices["hip"] = c_idx
                            break

            # Step 5: Generic circumference fallback only if field still missing
            if "waist" not in temp_indices or "hip" not in temp_indices:
                circ_cols = [
                    c_idx for c_idx, cell_val in enumerate(row_vals)
                    if cell_val is not None and ("circumfer" in str(cell_val).strip().lower() or "circ" in str(cell_val).strip().lower())
                    and c_idx != temp_indices.get("waist") and c_idx != temp_indices.get("hip")
                ]
                if "waist" not in temp_indices and circ_cols:
                    temp_indices["waist"] = circ_cols.pop(0)
                if "hip" not in temp_indices and circ_cols:
                    temp_indices["hip"] = circ_cols.pop(0)

            # Step 6: Guarantee waist and hip never share the same column index
            if "waist" in temp_indices and "hip" in temp_indices and temp_indices["waist"] == temp_indices["hip"]:
                del temp_indices["hip"]

            # We MUST have at least 'name' and 'weight' and 'employee_id'
            if "name" in temp_indices and "weight" in temp_indices and "employee_id" in temp_indices:
                header_row_idx = r_idx
                header_indices = temp_indices
                break

        if not header_row_idx:
            return {
                "success": False,
                "file_error": "Could not find required table columns ('Name' and 'Weight') in sheet headers."
            }

        eligible_data = {}
        incomplete_participants = []
        processed_ids = set()
        total_rows_processed = 0

        for row_idx in range(header_row_idx + 1, ws.max_row + 1):
            row_cells = [ws.cell(row=row_idx, column=col_idx + 1).value for col_idx in range(ws.max_column)]
            
            # 1. Ignore completely blank rows
            if all(val is None or str(val).strip() == "" for val in row_cells):
                continue

            total_rows_processed += 1

            raw_id = row_cells[header_indices["employee_id"]] if header_indices["employee_id"] < len(row_cells) else None
            raw_name = row_cells[header_indices["name"]] if header_indices["name"] < len(row_cells) else None
            raw_weight = row_cells[header_indices["weight"]] if header_indices["weight"] < len(row_cells) else None

            emp_id = cls.clean_employee_id(raw_id)
            name = cls.clean_name(raw_name)
            participant_display = name or (f"ID: {emp_id}" if emp_id else f"Row {row_idx}")

            missing_fields = []

            if not emp_id:
                missing_fields.append("Participant Identifier")
            elif emp_id in processed_ids:
                missing_fields.append("Duplicate Identifier")
            else:
                processed_ids.add(emp_id)

            # Weight validation
            weight = None
            if raw_weight is not None and str(raw_weight).strip() != "":
                try:
                    f_val = float(str(raw_weight).strip())
                    if f_val > 0:
                        weight = f_val
                except ValueError:
                    pass
            if weight is None:
                missing_fields.append("Weight")

            # Waist validation
            waist = None
            if "waist" in header_indices:
                raw_waist = row_cells[header_indices["waist"]] if header_indices["waist"] < len(row_cells) else None
                if raw_waist is not None and str(raw_waist).strip() != "":
                    try:
                        w_val = float(str(raw_waist).strip())
                        if w_val > 0:
                            waist = w_val
                    except ValueError:
                        pass
                if waist is None:
                    missing_fields.append("Waist")

            # Hip validation
            hip = None
            if "hip" in header_indices:
                raw_hip = row_cells[header_indices["hip"]] if header_indices["hip"] < len(row_cells) else None
                if raw_hip is not None and str(raw_hip).strip() != "":
                    try:
                        h_val = float(str(raw_hip).strip())
                        if h_val > 0:
                            hip = h_val
                    except ValueError:
                        pass
                if hip is None:
                    missing_fields.append("Hip")

            if missing_fields:
                incomplete_participants.append({
                    "row": row_idx,
                    "participant": participant_display,
                    "stage": sheet_label,
                    "missing_data": ", ".join(missing_fields),
                    "status": "Incomplete",
                    "reason": f"Missing {sheet_label} {', '.join(missing_fields)}"
                })
                continue

            whr = (waist / hip) if (waist and hip) else None

            eligible_data[emp_id] = {
                "name": name or emp_id,
                "weight": weight,
                "waist": waist,
                "hip": hip,
                "whr": whr
            }

        return {
            "success": True,
            "sheet_label": sheet_label,
            "total_rows": total_rows_processed,
            "valid_count": len(eligible_data),
            "incomplete_count": len(incomplete_participants),
            "data": eligible_data,
            "incomplete_participants": incomplete_participants,
            "file_error": None
        }

    @classmethod
    def calculate_results(cls, d0_res: dict, d45_res: dict, d90_res: dict) -> tuple[dict, list[dict], list[dict], list[str]]:
        """
        Merges datasets and calculates final rankings for eligible participants only.
        Classifies non-eligible participants into incomplete_records.
        Returns:
            summary: dict
            rankings: list of dicts (eligible only)
            incomplete_records: list of dicts (excluded participants)
            warnings: list of strings
        """
        d0_data = d0_res.get("data", {}) if isinstance(d0_res, dict) else d0_res
        d90_data = d90_res.get("data", {}) if isinstance(d90_res, dict) else d90_res
        d45_data = d45_res.get("data", {}) if isinstance(d45_res, dict) else d45_res

        # Collect all incomplete participant records from all 3 files
        incomplete_records = []
        if isinstance(d0_res, dict):
            incomplete_records.extend(d0_res.get("incomplete_participants", []))
        if isinstance(d45_res, dict):
            incomplete_records.extend(d45_res.get("incomplete_participants", []))
        if isinstance(d90_res, dict):
            incomplete_records.extend(d90_res.get("incomplete_participants", []))

        day0_ids = set(d0_data.keys())
        day90_ids = set(d90_data.keys())
        day45_ids = set(d45_data.keys())

        # Participants in Day 0 but missing in Day 90
        only_in_d0 = day0_ids - day90_ids
        for eid in sorted(only_in_d0):
            incomplete_records.append({
                "row": "-",
                "participant": d0_data[eid].get("name", eid),
                "stage": "Day 90",
                "missing_data": "Final Measurements (Day 90)",
                "status": "Incomplete",
                "reason": f"Participant '{d0_data[eid].get('name', eid)}' ({eid}) has Day 0 measurements but is missing in Day 90 sheet."
            })

        # Participants in Day 90 but missing in Day 0
        only_in_d90 = day90_ids - day0_ids
        for eid in sorted(only_in_d90):
            incomplete_records.append({
                "row": "-",
                "participant": d90_data[eid].get("name", eid),
                "stage": "Day 0",
                "missing_data": "Baseline Measurements (Day 0)",
                "status": "Incomplete",
                "reason": f"Participant '{d90_data[eid].get('name', eid)}' ({eid}) has Day 90 measurements but is missing in Day 0 sheet."
            })

        common_ids = day0_ids & day90_ids

        warnings = []
        raw_results = []
        max_weight_loss_pct = 0.0
        max_whr_imp_pct = 0.0
        has_whr_data = False

        for eid in sorted(common_ids):
            try:
                d0 = d0_data[eid]
                d90 = d90_data[eid]

                d0_wt = d0.get("weight")
                d90_wt = d90.get("weight")
                d0_waist = d0.get("waist")
                d90_waist = d90.get("waist")
                d0_hip = d0.get("hip")
                d90_hip = d90.get("hip")

                d0_whr = d0.get("whr")
                d90_whr = d90.get("whr")

                # Compute WHR dynamically if waist & hip are valid numbers
                if d0_whr is None and d0_waist and d0_hip and d0_hip > 0:
                    d0_whr = d0_waist / d0_hip
                if d90_whr is None and d90_waist and d90_hip and d90_hip > 0:
                    d90_whr = d90_waist / d90_hip

                missing_items = []
                if d0_wt is None or d0_wt <= 0: missing_items.append("Day 0 Weight")
                if d90_wt is None or d90_wt <= 0: missing_items.append("Day 90 Weight")
                if d0_waist is None or d0_waist <= 0: missing_items.append("Day 0 Waist")
                if d90_waist is None or d90_waist <= 0: missing_items.append("Day 90 Waist")
                if d0_hip is None or d0_hip <= 0: missing_items.append("Day 0 Hip")
                if d90_hip is None or d90_hip <= 0: missing_items.append("Day 90 Hip")

                if missing_items:
                    incomplete_records.append({
                        "row": "-",
                        "participant": d0.get("name", eid),
                        "stage": "Day 0 / Day 90",
                        "missing_data": ", ".join(missing_items),
                        "status": "Incomplete",
                        "reason": f"Participant '{d0.get('name', eid)}' ({eid}) missing required measurements: {', '.join(missing_items)}"
                    })
                    continue

                # Weight Loss % = ((Day0Weight - Day90Weight) / Day0Weight) * 100
                weight_loss_pct = ((d0_wt - d90_wt) / d0_wt) * 100.0

                # WHR Improvement % = ((Day0WHR - Day90WHR) / Day0WHR) * 100
                whr_imp_pct = 0.0
                if d0_whr is not None and d90_whr is not None and d0_whr > 0:
                    whr_imp_pct = ((d0_whr - d90_whr) / d0_whr) * 100.0
                    has_whr_data = True

                if weight_loss_pct > max_weight_loss_pct:
                    max_weight_loss_pct = weight_loss_pct
                if whr_imp_pct > max_whr_imp_pct:
                    max_whr_imp_pct = whr_imp_pct

                if eid not in day45_ids:
                    warnings.append(f"Participant '{d0.get('name', eid)}' ({eid}) is missing in Day 45 (included in final rankings).")

                raw_results.append({
                    "employee_id": eid,
                    "name": d0.get("name", eid),
                    "day0_weight": float(d0_wt),
                    "day90_weight": float(d90_wt),
                    "weight_loss_percent": float(weight_loss_pct),
                    "day0_whr": float(d0_whr) if d0_whr is not None else None,
                    "day90_whr": float(d90_whr) if d90_whr is not None else None,
                    "whr_improvement_percent": float(whr_imp_pct)
                })

            except Exception as p_err:
                incomplete_records.append({
                    "row": "-",
                    "participant": d0_data.get(eid, {}).get("name", eid),
                    "stage": "Calculation Error",
                    "missing_data": str(p_err),
                    "status": "Incomplete",
                    "reason": f"Calculation error for participant '{eid}': {str(p_err)}"
                })
                continue

        if not raw_results:
            raise ValueError("No eligible participants with complete Day 0 and Day 90 measurements found.")

        norm_weight_denom = max_weight_loss_pct if max_weight_loss_pct > 0 else 1.0
        norm_whr_denom = max_whr_imp_pct if max_whr_imp_pct > 0 else 1.0

        for r in raw_results:
            r["normalized_weight"] = max(0.0, (r["weight_loss_percent"] / norm_weight_denom) * 100.0) if max_weight_loss_pct > 0 else 0.0
            r["normalized_whr"] = max(0.0, (r["whr_improvement_percent"] / norm_whr_denom) * 100.0) if max_whr_imp_pct > 0 else 0.0
            
            if has_whr_data:
                r["final_score"] = (r["normalized_weight"] * 0.5) + (r["normalized_whr"] * 0.5)
            else:
                r["final_score"] = r["normalized_weight"]

        sorted_results = sorted(
            raw_results, 
            key=lambda x: (x["final_score"], x["weight_loss_percent"], x["whr_improvement_percent"]), 
            reverse=True
        )

        for idx, r in enumerate(sorted_results, 1):
            r["rank"] = idx

        winner = sorted_results[0] if sorted_results else None

        total_d0 = d0_res.get("total_rows", len(day0_ids)) if isinstance(d0_res, dict) else len(day0_ids)
        total_d90 = d90_res.get("total_rows", len(day90_ids)) if isinstance(d90_res, dict) else len(day90_ids)

        summary = {
            "total_processed_rows": total_d0 + total_d90,
            "total_participants": len(sorted_results),
            "eligible_count": len(sorted_results),
            "incomplete_count": len(incomplete_records),
            "highest_weight_loss": round(max_weight_loss_pct, 2),
            "highest_whr_improvement": round(max_whr_imp_pct, 2) if has_whr_data else 0.0,
            "winner_name": winner["name"] if winner else "N/A",
            "winner_employee_id": winner["employee_id"] if winner else "N/A",
            "winner_final_score": round(winner["final_score"], 2) if winner else 0.0
        }

        for r in sorted_results:
            r["day0_weight"] = round(r["day0_weight"], 2) if r["day0_weight"] is not None else 0.0
            r["day90_weight"] = round(r["day90_weight"], 2) if r["day90_weight"] is not None else 0.0
            r["weight_loss_percent"] = round(r["weight_loss_percent"], 2) if r["weight_loss_percent"] is not None else 0.0
            r["day0_whr"] = round(r["day0_whr"], 4) if r["day0_whr"] is not None else "-"
            r["day90_whr"] = round(r["day90_whr"], 4) if r["day90_whr"] is not None else "-"
            r["whr_improvement_percent"] = round(r["whr_improvement_percent"], 2) if r["whr_improvement_percent"] is not None else 0.0
            r["normalized_weight"] = round(r["normalized_weight"], 2) if r["normalized_weight"] is not None else 0.0
            r["normalized_whr"] = round(r["normalized_whr"], 2) if r["normalized_whr"] is not None else 0.0
            r["final_score"] = round(r["final_score"], 2) if r["final_score"] is not None else 0.0

        return summary, sorted_results, incomplete_records, warnings

    @classmethod
    def generate_excel_bytes(cls, rankings: list[dict], title: str) -> bytes:
        """Generates a styled Excel sheet using openpyxl and returns bytes."""
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:31]

        header_fill = PatternFill("solid", fgColor="4F46E5")
        header_font = Font(bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        headers = [
            "Rank", "Identifier", "Name", 
            "Day 0 Weight", "Day 90 Weight", "Weight Loss %", 
            "Day 0 WHR", "Day 90 WHR", "WHR Improvement %", 
            "Final Score"
        ]

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        for row_idx, r in enumerate(rankings, 2):
            ws.cell(row=row_idx, column=1, value=r["rank"]).alignment = center_align
            ws.cell(row=row_idx, column=2, value=r["employee_id"]).alignment = center_align
            ws.cell(row=row_idx, column=3, value=r["name"]).alignment = left_align
            ws.cell(row=row_idx, column=4, value=r["day0_weight"]).alignment = right_align
            ws.cell(row=row_idx, column=5, value=r["day90_weight"]).alignment = right_align
            ws.cell(row=row_idx, column=6, value=r["weight_loss_percent"]).alignment = right_align
            ws.cell(row=row_idx, column=7, value=r["day0_whr"]).alignment = right_align
            ws.cell(row=row_idx, column=8, value=r["day90_whr"]).alignment = right_align
            ws.cell(row=row_idx, column=9, value=r["whr_improvement_percent"]).alignment = right_align
            ws.cell(row=row_idx, column=10, value=r["final_score"]).alignment = right_align

            for col_idx in range(1, 11):
                ws.cell(row=row_idx, column=col_idx).border = thin_border

        ws.row_dimensions[1].height = 25
        for row in range(2, len(rankings) + 2):
            ws.row_dimensions[row].height = 20

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(14, max_len + 3)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()
